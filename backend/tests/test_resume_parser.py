import io
import zipfile

import pytest
from docx import Document

from app.modules.resume.parser import (
    _docx_xml_text,
    parse_docx_bytes,
    parse_resume_bytes,
)


def _docx_with_xml_text(text: str) -> bytes:
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:body><w:p><w:r><w:t>"
        + text
        + "</w:t></w:r></w:p></w:body></w:document>"
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


def test_docx_xml_fallback_extracts_text():
    data = _docx_with_xml_text("文本框里的技能")
    assert "文本框里的技能" in _docx_xml_text(data)


def test_parse_docx_bytes_reads_paragraph_and_table():
    doc = Document()
    doc.add_paragraph("河北环境工程学院")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "技能"
    table.rows[0].cells[1].text = "Python"
    buffer = io.BytesIO()
    doc.save(buffer)

    text = parse_docx_bytes(buffer.getvalue())
    assert "河北环境工程学院" in text
    assert "Python" in text


def test_parse_resume_bytes_empty_docx_raises():
    doc = Document()
    buffer = io.BytesIO()
    doc.save(buffer)
    with pytest.raises(ValueError):
        parse_resume_bytes("empty.docx", buffer.getvalue())
