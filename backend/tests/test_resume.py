import pytest
from io import BytesIO
from docx import Document

from app.models.resume import Resume


def test_resume_model_creation(db_session):
    """测试简历模型创建"""
    resume = Resume(
        user_id="test_user",
        version=1,
        raw_text="Test resume content",
        source_file="test.pdf",
    )
    db_session.add(resume)
    db_session.commit()
    assert resume.id is not None
    assert resume.version == 1


def test_resume_list_empty(db_session):
    """测试空简历列表"""
    resumes = db_session.query(Resume).all()
    assert len(resumes) == 0


def test_resume_versions_api(client):
    """测试简历版本历史接口"""
    first = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": "张三，3年Python开发经验，熟悉Django", "source_file": "a.txt"},
    )
    assert first.json()["success"] is True
    resume_id = first.json()["data"]["id"]

    client.post(
        "/api/v1/resume/upload",
        json={"raw_text": "张三，5年Python开发经验，熟悉Django和FastAPI", "source_file": "b.txt"},
    )

    versions = client.get(f"/api/v1/resume/{resume_id}/versions")
    assert versions.status_code == 200
    data = versions.json()["data"]
    assert len(data) == 2
    assert data[0]["version"] == 2
    assert data[0]["parsed_json"]["skills"]


def test_resume_upload_rejects_empty_text(client):
    """测试空文本简历不能上传"""
    response = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": "   ", "source_file": "empty.txt"},
    )
    assert response.status_code == 400


def test_resume_upload_docx(client):
    """测试 .docx 简历上传"""
    doc = Document()
    doc.add_paragraph("张三，5年Python开发经验，熟悉Django和FastAPI")
    buffer = BytesIO()
    doc.save(buffer)

    response = client.post(
        "/api/v1/resume/upload-file",
        files={
            "file": (
                "resume.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 200
    data = response.json()["data"]
    assert data["filename"] == "resume.docx"

    versions = client.get(f"/api/v1/resume/{data['id']}/versions")
    assert versions.status_code == 200
    assert "Django" in versions.json()["data"][0]["raw_text"]


def test_resume_upload_docx_empty_rejected(client):
    """测试内容为空的 docx 上传会被拒绝"""
    doc = Document()
    buffer = BytesIO()
    doc.save(buffer)

    response = client.post(
        "/api/v1/resume/upload-file",
        files={
            "file": (
                "empty.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 400


def test_resume_export_pdf_and_word(client):
    """测试简历导出 PDF / Word"""
    upload = client.post(
        "/api/v1/resume/upload",
        json={
            "raw_text": "John Doe\nEducation: CS Master\nProject: Payment System\nSkills: Python, Docker",
            "source_file": "a.txt",
        },
    )
    resume_id = upload.json()["data"]["id"]

    pdf = client.get(f"/api/v1/resume/{resume_id}/export?format=pdf")
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert len(pdf.content) > 100

    word = client.get(f"/api/v1/resume/{resume_id}/export?format=word")
    assert word.status_code == 200
    assert "wordprocessingml" in word.headers["content-type"]
    assert len(word.content) > 100


def test_resume_optimize_then_export(client, sample_resume_text, sample_jd_text):
    """测试优化后生成新版本并可直接导出，且不丢失原文"""
    upload = client.post(
        "/api/v1/resume/upload",
        json={"raw_text": sample_resume_text, "source_file": "a.txt"},
    )
    resume_id = upload.json()["data"]["id"]

    optimize = client.post(
        "/api/v1/resume/optimize",
        json={"resume_id": resume_id, "jd_text": sample_jd_text},
    )
    assert optimize.status_code == 200
    data = optimize.json()["data"]
    assert data["optimized"].strip() == sample_resume_text
    assert data["preservation"]["passed"] is True
    assert data["new_version"] is not None

    new_version_id = data["new_version"]["id"]
    word = client.get(f"/api/v1/resume/{new_version_id}/export?format=word")
    assert word.status_code == 200
    assert "wordprocessingml" in word.headers["content-type"]
    assert len(word.content) > 100
