"""简历文件解析模块。

只负责把上传文件或文本转换成纯文本，不依赖数据库或其他业务模块。
"""

import io
import re
import zipfile
from typing import List
from xml.etree import ElementTree

import markdown
import PyPDF2
from bs4 import BeautifulSoup
from docx import Document


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W_T = f"{{{W_NS}}}t"
W_P = f"{{{W_NS}}}p"

EMPTY_RESUME_ERROR = "未能从文件中识别到简历内容，请上传可编辑文本的 PDF/Word/Markdown/TXT，或确认截图清晰"


def parse_text(raw_text: str) -> str:
    """清理纯文本。"""
    text = re.sub(r"\n\s*\n", "\n\n", raw_text or "")
    return text.strip()


def parse_markdown_text(content: str) -> str:
    html = markdown.markdown(content or "")
    soup = BeautifulSoup(html, "html.parser")
    return parse_text(soup.get_text(separator="\n"))


def parse_markdown(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        return parse_markdown_text(f.read())


def parse_pdf_bytes(file_bytes: bytes) -> str:
    try:
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        text = "".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        raise ValueError(f"PDF 解析失败: {e}")
    return parse_text(text)


def parse_pdf(file_path: str) -> str:
    with open(file_path, "rb") as f:
        return parse_pdf_bytes(f.read())


def _docx_xml_text(file_bytes: bytes) -> str:
    """解析 docx 中所有 XML 文本，覆盖文本框、页眉页脚和脚注。"""
    try:
        archive = zipfile.ZipFile(io.BytesIO(file_bytes))
    except zipfile.BadZipFile:
        return ""

    xml_names = [
        name for name in archive.namelist()
        if name.startswith("word/") and name.endswith(".xml")
    ]

    def part_priority(name: str) -> int:
        if name == "word/document.xml":
            return 0
        if "header" in name:
            return 1
        if "footer" in name:
            return 2
        return 3

    xml_names.sort(key=part_priority)
    lines: List[str] = []
    for name in xml_names:
        try:
            root = ElementTree.fromstring(archive.read(name))
        except Exception:
            continue
        for paragraph in root.iter(W_P):
            text = "".join((node.text or "") for node in paragraph.iter(W_T))
            text = re.sub(r"\s+", " ", text).strip()
            if text:
                lines.append(text)
    return "\n".join(lines)


def _docx_images(file_bytes: bytes) -> List[bytes]:
    try:
        archive = zipfile.ZipFile(io.BytesIO(file_bytes))
    except zipfile.BadZipFile:
        return []
    return [
        archive.read(name)
        for name in archive.namelist()
        if name.startswith("word/media/")
        and name.lower().endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp"))
    ]


def _ocr_images(images: List[bytes]) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""

    parts: List[str] = []
    for data in images:
        try:
            text = pytesseract.image_to_string(
                Image.open(io.BytesIO(data)), lang="chi_sim+eng"
            )
            if text and text.strip():
                parts.append(parse_text(text))
        except Exception:
            continue
    return "\n".join(parts)


def parse_docx_bytes(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
        lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    lines.append("\t".join(cells))
        for section in document.sections:
            headers = [section.header, section.first_page_header, section.even_page_header]
            footers = [section.footer, section.first_page_footer, section.even_page_footer]
            for header in headers:
                if header is not None:
                    lines.extend(p.text.strip() for p in header.paragraphs if p.text.strip())
            for footer in footers:
                if footer is not None:
                    lines.extend(p.text.strip() for p in footer.paragraphs if p.text.strip())
        text = parse_text("\n".join(lines))
    except Exception as e:
        raise ValueError(f"DOCX 解析失败: {e}")

    if not text:
        text = parse_text(_docx_xml_text(file_bytes))
    if not text:
        text = _ocr_images(_docx_images(file_bytes))
    return text


def parse_resume_file(file_path: str) -> str:
    name = (file_path or "").lower()
    if name.endswith(".pdf"):
        return parse_pdf(file_path)
    if name.endswith(".md") or name.endswith(".markdown"):
        return parse_markdown(file_path)
    if name.endswith(".txt"):
        with open(file_path, "r", encoding="utf-8") as f:
            return parse_text(f.read())
    if name.endswith(".docx"):
        with open(file_path, "rb") as f:
            return parse_docx_bytes(f.read())
    raise ValueError(f"不支持的文件类型: {file_path}")


def parse_resume_bytes(filename: str, file_bytes: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = parse_pdf_bytes(file_bytes)
    elif name.endswith(".md") or name.endswith(".markdown"):
        text = parse_markdown_text(file_bytes.decode("utf-8"))
    elif name.endswith(".txt"):
        text = parse_text(file_bytes.decode("utf-8"))
    elif name.endswith(".docx"):
        text = parse_docx_bytes(file_bytes)
    else:
        raise ValueError(f"不支持的文件类型: {filename}")

    if not text:
        raise ValueError(EMPTY_RESUME_ERROR)
    return text
